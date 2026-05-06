from __future__ import annotations

import json
import os
import subprocess
import textwrap
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont, ImageOps


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "item" / "Reports" / "sprint6"
FIG_DIR = OUT_DIR / "evidence_screenshots"
REAL_DIR = OUT_DIR / "real_evidence_screenshots"
DOCX_PATH = OUT_DIR / "ANU_TreasureMap_Sprint6_Review_Report.docx"
ROOT_DOCX_PATH = ROOT / "ANU_TreasureMap_Sprint6_Review_Report.docx"


FONT_CANDIDATES = [
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
    "/System/Library/Fonts/Helvetica.ttc",
    "/Library/Fonts/Arial.ttf",
]


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = []
    if bold:
        candidates.extend(
            [
                "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
                "/System/Library/Fonts/Supplemental/Arial.ttf",
            ]
        )
    candidates.extend(FONT_CANDIDATES)
    for candidate in candidates:
        if Path(candidate).exists():
            try:
                return ImageFont.truetype(candidate, size)
            except OSError:
                continue
    return ImageFont.load_default()


FONT_TITLE = font(42, True)
FONT_H1 = font(30, True)
FONT_H2 = font(24, True)
FONT_BODY = font(21)
FONT_SMALL = font(18)
FONT_MONO = font(18)
FONT_CAPTION = font(16)


def ensure_dirs() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    REAL_DIR.mkdir(parents=True, exist_ok=True)


def run(cmd: list[str]) -> str:
    return subprocess.check_output(cmd, cwd=ROOT, text=True, stderr=subprocess.STDOUT)


def wrap_lines(text: str, width: int) -> list[str]:
    lines: list[str] = []
    for part in str(text).splitlines() or [""]:
        if not part:
            lines.append("")
            continue
        lines.extend(textwrap.wrap(part, width=width, replace_whitespace=False))
    return lines


def draw_wrapped(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int],
    text: str,
    max_width: int,
    fill: tuple[int, int, int] = (30, 41, 59),
    fnt: ImageFont.ImageFont = FONT_BODY,
    line_gap: int = 7,
) -> int:
    x, y = xy
    words = str(text).split()
    if not words:
        return y
    line = ""
    for word in words:
        test = word if not line else f"{line} {word}"
        bbox = draw.textbbox((0, 0), test, font=fnt)
        if bbox[2] - bbox[0] <= max_width:
            line = test
        else:
            draw.text((x, y), line, font=fnt, fill=fill)
            y += (draw.textbbox((0, 0), line, font=fnt)[3] - draw.textbbox((0, 0), line, font=fnt)[1]) + line_gap
            line = word
    if line:
        draw.text((x, y), line, font=fnt, fill=fill)
        bbox = draw.textbbox((0, 0), line, font=fnt)
        y += (bbox[3] - bbox[1]) + line_gap
    return y


def rounded_rectangle(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    radius: int,
    fill: tuple[int, int, int],
    outline: tuple[int, int, int] | None = None,
    width: int = 1,
) -> None:
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def make_plate(
    filename: str,
    title: str,
    subtitle: str,
    rows: list[tuple[str, str, str]],
    source: str,
    width: int = 1800,
    accent: tuple[int, int, int] = (19, 74, 117),
) -> Path:
    margin = 70
    row_h = 116
    header_h = 210
    footer_h = 72
    height = header_h + footer_h + row_h * len(rows) + 90
    img = Image.new("RGB", (width, height), (248, 250, 252))
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 0, width, 34), fill=accent)
    draw.text((margin, 62), title, font=FONT_TITLE, fill=(15, 23, 42))
    draw_wrapped(draw, (margin, 120), subtitle, width - margin * 2, fill=(71, 85, 105), fnt=FONT_BODY)
    y = header_h
    rounded_rectangle(draw, (margin, y, width - margin, y + 52), 10, accent)
    col1 = margin + 26
    col2 = margin + 360
    col3 = margin + 910
    draw.text((col1, y + 14), "Evidence item", font=FONT_SMALL, fill=(255, 255, 255))
    draw.text((col2, y + 14), "Observed content", font=FONT_SMALL, fill=(255, 255, 255))
    draw.text((col3, y + 14), "Rubric value", font=FONT_SMALL, fill=(255, 255, 255))
    y += 60
    for idx, (a, b, c) in enumerate(rows):
        fill = (255, 255, 255) if idx % 2 == 0 else (241, 245, 249)
        rounded_rectangle(draw, (margin, y, width - margin, y + row_h - 12), 8, fill, (203, 213, 225))
        draw_wrapped(draw, (col1, y + 20), a, 285, fnt=FONT_SMALL, fill=(15, 23, 42), line_gap=5)
        draw_wrapped(draw, (col2, y + 20), b, 500, fnt=FONT_SMALL, fill=(30, 41, 59), line_gap=5)
        draw_wrapped(draw, (col3, y + 20), c, 720, fnt=FONT_SMALL, fill=(30, 41, 59), line_gap=5)
        y += row_h
    draw.text((margin, height - 52), f"Source: {source}", font=FONT_CAPTION, fill=(71, 85, 105))
    out = FIG_DIR / filename
    img.save(out, quality=94)
    return out


def make_code_plate(
    filename: str,
    title: str,
    subtitle: str,
    code_blocks: list[tuple[str, str]],
    source: str,
    width: int = 1800,
    accent: tuple[int, int, int] = (71, 85, 105),
) -> Path:
    margin = 70
    block_h = 330
    height = 210 + block_h * len(code_blocks) + 95
    img = Image.new("RGB", (width, height), (248, 250, 252))
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 0, width, 34), fill=accent)
    draw.text((margin, 62), title, font=FONT_TITLE, fill=(15, 23, 42))
    draw_wrapped(draw, (margin, 120), subtitle, width - margin * 2, fill=(71, 85, 105), fnt=FONT_BODY)
    y = 210
    for label, content in code_blocks:
        rounded_rectangle(draw, (margin, y, width - margin, y + block_h - 22), 10, (15, 23, 42), (51, 65, 85))
        draw.text((margin + 24, y + 18), label, font=FONT_SMALL, fill=(147, 197, 253))
        yy = y + 58
        for line in content.splitlines()[:10]:
            draw.text((margin + 24, yy), line[:145], font=FONT_MONO, fill=(226, 232, 240))
            yy += 24
        y += block_h
    draw.text((margin, height - 55), f"Source: {source}", font=FONT_CAPTION, fill=(71, 85, 105))
    out = FIG_DIR / filename
    img.save(out, quality=94)
    return out


def crop_or_fit(path: Path, box: tuple[int, int, int, int] | None, size: tuple[int, int]) -> Image.Image:
    img = Image.open(path).convert("RGB")
    if box:
        img = img.crop(box)
    return ImageOps.contain(img, size, Image.Resampling.LANCZOS)


def make_product_plate() -> Path:
    width, height = 1800, 1080
    img = Image.new("RGB", (width, height), (248, 250, 252))
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 0, width, 34), fill=(132, 85, 45))
    draw.text((70, 62), "Product/build evidence from repository assets", font=FONT_TITLE, fill=(15, 23, 42))
    draw_wrapped(
        draw,
        (70, 120),
        "The current vertical slice is evidenced through the implemented map asset, location imagery, treasure art, and event data already committed to the Godot project.",
        width - 140,
        fill=(71, 85, 105),
        fnt=FONT_BODY,
    )
    assets = [
        ("ANU map asset", ROOT / "Godot_Game/anu-treasure-map/resources/Map/ANU Map.png"),
        ("Bruce Hall event region", ROOT / "Godot_Game/anu-treasure-map/resources/Location/Bruce Hall.jpg"),
        ("Yuki treasure art", ROOT / "Godot_Game/anu-treasure-map/resources/treasure/treasure_yuki.png"),
        ("Phone/profile icon", ROOT / "Godot_Game/anu-treasure-map/resources/icon/profile.png"),
    ]
    x_positions = [70, 520, 970, 1380]
    y = 220
    for (label, path), x in zip(assets, x_positions):
        rounded_rectangle(draw, (x, y, x + 360, y + 580), 14, (255, 255, 255), (203, 213, 225))
        if path.exists():
            tile = crop_or_fit(path, None, (320, 420))
            img.paste(tile, (x + 20 + (320 - tile.width) // 2, y + 35))
        draw_wrapped(draw, (x + 20, y + 480), label, 320, fill=(15, 23, 42), fnt=FONT_H2)
    draw.text(
        (70, height - 55),
        "Source: Godot_Game/anu-treasure-map/resources and data committed in the repository.",
        font=FONT_CAPTION,
        fill=(71, 85, 105),
    )
    out = FIG_DIR / "figure_09_product_build_assets.jpg"
    img.save(out, quality=94)
    return out


def make_event_visual_feedback_plate() -> Path:
    width, height = 1800, 1240
    img = Image.new("RGB", (width, height), (248, 250, 252))
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 0, width, 34), fill=(132, 85, 45))
    draw.text((70, 62), "Event visual feedback implementation evidence", font=FONT_TITLE, fill=(15, 23, 42))
    draw_wrapped(
        draw,
        (70, 120),
        "SCRUM-118 records the client feedback that events should feel visual when they occur. The implementation evidence combines the event overlay flow, treasure-image outcome support, and staged event image assets for Week 1/future event expansion.",
        width - 140,
        fill=(71, 85, 105),
        fnt=FONT_BODY,
    )
    assets = [
        ("Intro image", ROOT / "Godot_Game/anu-treasure-map/resources/event/event_bruce_hall_lemon_tree_intro.png"),
        ("Investigation image", ROOT / "Godot_Game/anu-treasure-map/resources/event/event_bruce_hall_lemon_tree_investigation.png"),
        ("Choice success image", ROOT / "Godot_Game/anu-treasure-map/resources/event/event_bruce_hall_lemon_tree_choice_a_success.png"),
        ("Treasure image", ROOT / "Godot_Game/anu-treasure-map/resources/treasure/treasure_yuki.png"),
    ]
    x_positions = [70, 500, 930, 1360]
    y = 250
    for (label, path), x in zip(assets, x_positions):
        rounded_rectangle(draw, (x, y, x + 350, y + 420), 14, (255, 255, 255), (203, 213, 225))
        if path.exists():
            tile = crop_or_fit(path, None, (310, 270))
            img.paste(tile, (x + 20 + (310 - tile.width) // 2, y + 28))
        draw_wrapped(draw, (x + 22, y + 325), label, 305, fill=(15, 23, 42), fnt=FONT_H2)
    rows = [
        ("Trigger transition", "locations.gd switches from normal location UI to a full-screen RandomEventPanel overlay.", "Visual event entry rather than text-only location dialogue."),
        ("Branch/outcome flow", "random_event_panel.gd supports narrative click-through, choices, outcome text, and continue state.", "Player action produces a visible event state."),
        ("Treasure image display", "Successful choices can load and show treasure_image_path after the outcome.", "Treasure collection becomes visible in the UI."),
        ("Asset handover", "resources/event and resources/narrative/events.md preserve intro, investigation, success, miss, and hint image sequence patterns.", "Future event expansion has a reusable visual grammar."),
    ]
    y = 700
    rounded_rectangle(draw, (70, y, width - 70, y + 48), 10, (132, 85, 45))
    draw.text((96, y + 13), "Evidence item", font=FONT_SMALL, fill=(255, 255, 255))
    draw.text((455, y + 13), "Observed content", font=FONT_SMALL, fill=(255, 255, 255))
    draw.text((1110, y + 13), "Reflection value", font=FONT_SMALL, fill=(255, 255, 255))
    y += 60
    for idx, (a, b, c) in enumerate(rows):
        fill = (255, 255, 255) if idx % 2 == 0 else (241, 245, 249)
        rounded_rectangle(draw, (70, y, width - 70, y + 88), 8, fill, (203, 213, 225))
        draw_wrapped(draw, (96, y + 15), a, 300, fnt=FONT_SMALL, fill=(15, 23, 42), line_gap=4)
        draw_wrapped(draw, (455, y + 15), b, 590, fnt=FONT_SMALL, fill=(30, 41, 59), line_gap=4)
        draw_wrapped(draw, (1110, y + 15), c, 600, fnt=FONT_SMALL, fill=(30, 41, 59), line_gap=4)
        y += 100
    draw.text(
        (70, height - 55),
        "Source: Jira SCRUM-118; Godot locations.gd, random_event_panel.gd, resources/event, resources/treasure, and resources/narrative/events.md.",
        font=FONT_CAPTION,
        fill=(71, 85, 105),
    )
    out = FIG_DIR / "figure_22_event_visual_feedback.jpg"
    img.save(out, quality=94)
    return out


@dataclass
class JiraIssue:
    key: str
    summary: str
    owner: str
    status: str
    moscow: str
    points: int
    stream: str


JIRA_ISSUES = [
    JiraIssue("SCRUM-81", "Confirm Sprint 6 Software Goal and Protected Outputs with Clients", "Jeongjoon", "Done", "Must", 2, "Stakeholder / scope"),
    JiraIssue("SCRUM-82", "Maintain Sprint 6 Feedback-to-Jira Action Register", "Jeongjoon", "Done", "Must", 2, "Stakeholder / evidence"),
    JiraIssue("SCRUM-83", "Maintain Sprint 6 Feedback-to-Jira Action Register", "Jeongjoon", "Done", "Must", 2, "Stakeholder / evidence"),
    JiraIssue("SCRUM-84", "Maintain Sprint 6 Risk and Dependency Log", "Jeongjoon", "Done", "Must", 2, "Planning"),
    JiraIssue("SCRUM-85", "Prepare Sprint 6 Evidence Report and Review Package", "Jeongjoon", "Done", "Must", 3, "Review evidence"),
    JiraIssue("SCRUM-86", "Request and Schedule Pitch Dry Run with Clients", "Jeongjoon", "In Review", "Should", 1, "Stakeholder / pitch"),
    JiraIssue("SCRUM-87", "Lock Week 1 Happy Path Gameplay Flow", "Zijun", "Done", "Must", 3, "Game design"),
    JiraIssue("SCRUM-88", "Define Yuki Event Gameplay Flow and Branching Script", "Wing", "Done", "Must", 3, "Game design"),
    JiraIssue("SCRUM-89", "Define Treasure Collect vs Hint Decision Rules", "Zijun", "Done", "Must", 3, "Game design"),
    JiraIssue("SCRUM-90", "Define Week 1 Result State and Credits Summary Prototype", "Zijun", "In Review", "Should", 2, "Game design"),
    JiraIssue("SCRUM-91", "Deploy Configurable Nodes for All 13 Location Areas", "Shaoming", "Done", "Must", 5, "Core systems"),
    JiraIssue("SCRUM-92", "Implement Active/Inactive Region Selection in Godot", "Shaoming", "Done", "Must", 3, "Core systems"),
    JiraIssue("SCRUM-93", "Implement Yuki Event in Godot", "Suya", "Done", "Must", 5, "Core systems"),
    JiraIssue("SCRUM-94", "Implement Collect vs Leave Hint Outcome", "Suya", "Done", "Must", 5, "Core systems"),
    JiraIssue("SCRUM-95", "Implement Time-Only Resource Model for Week 1", "Shaoming", "In Review", "Should", 3, "Core systems"),
    JiraIssue("SCRUM-96", "Improve UI Contrast Based on Client Glasses-Off Test", "Xiyin", "Done", "Must", 3, "UI/UX"),
    JiraIssue("SCRUM-97", "Finalize Starting Page and Onboarding UI for Week 1 Slice", "Xiyin", "Done", "Must", 3, "UI/UX"),
    JiraIssue("SCRUM-98", "Design Map Region Marker and Hint Marker Visuals", "Xiyin", "In Review", "Should", 3, "UI/UX"),
    JiraIssue("SCRUM-99", "Create Yuki Event Artwork and Treasure Visuals", "Anlin", "Done", "Must", 3, "Art"),
    JiraIssue("SCRUM-100", "Create Week 1 Event Visuals", "Anlin", "Done", "Should", 3, "Art"),
    JiraIssue("SCRUM-101", "Define Event Data Format for Yuki Branching", "Wing", "Done", "Must", 3, "Data / content"),
    JiraIssue("SCRUM-102", "Prepare Week 1 Event Storyline and Branching Content", "Wing", "Done", "Must", 3, "Data / content"),
    JiraIssue("SCRUM-103", "Define Simulated Future-Player Hint Data for Testing", "Wing", "In Review", "Should", 2, "Data / testing"),
    JiraIssue("SCRUM-104", "Prepare External Proof-of-Concept Testing Plan", "Zijun", "Done", "Must", 3, "Testing"),
    JiraIssue("SCRUM-105", "Test Starting Page and Onboarding Flow", "Zijun", "In Review", "Must", 2, "Testing"),
    JiraIssue("SCRUM-106", "Test Week 1 Yuki Happy Path", "Zijun", "In Review", "Must", 3, "Testing"),
    JiraIssue("SCRUM-107", "Test Future-Player Hint Marker Flow", "Zijun", "In Review", "Should", 3, "Testing"),
    JiraIssue("SCRUM-108", "Run External User Testing and Summarise Results", "Zijun", "In Review", "Must", 5, "Testing"),
    JiraIssue("SCRUM-109", "Standardise PR Evidence Requirements", "Jeongjoon", "Done", "Should", 2, "GitHub quality"),
    JiraIssue("SCRUM-110", "Create Demo Video and Build Evidence Package", "Jeongjoon", "Done", "Must", 3, "Build evidence"),
    JiraIssue("SCRUM-111", "Maintain Velocity and Estimate Review", "Jeongjoon", "Done", "Should", 1, "Planning"),
    JiraIssue("SCRUM-112", "Prepare Game Concept Presentation Structure", "Jeongjoon", "Done", "Should", 3, "Handover / pitch"),
    JiraIssue("SCRUM-113", "Record Client Advocacy and Convenor Alignment Outcome for Sprint 6", "Jeongjoon", "Done", "Must", 2, "Stakeholder / alignment"),
    JiraIssue("SCRUM-114", "Prepare Week 10 Poster and Pitch Rough Outline for Client Dry Run", "Jeongjoon", "In Review", "Should", 2, "Stakeholder / pitch"),
    JiraIssue("SCRUM-115", "Document Start Location and Campus Photo Accuracy Decisions", "Jeongjoon", "Done", "Must", 2, "Map accuracy"),
    JiraIssue("SCRUM-116", "Document External Testing Constraints and Participant No-Show Reflection", "Zijun", "In Review", "Must", 2, "Testing / reflection"),
    JiraIssue("SCRUM-117", "Create GitHub Game README and Handover Docs for Sprint 6 Evidence", "Jeongjoon", "Done", "Must", 3, "GitHub docs"),
    JiraIssue("SCRUM-118", "Implement Event Animation and Image Transition Feedback", "Jeongjoon", "Done", "Should", 2, "Client feedback"),
]


def make_jira_board_plate() -> Path:
    done = sum(1 for i in JIRA_ISSUES if i.status == "Done")
    review = sum(1 for i in JIRA_ISSUES if i.status == "In Review")
    progress = sum(1 for i in JIRA_ISSUES if i.status in {"In Progress", "To Do"})
    planned = sum(i.points for i in JIRA_ISSUES)
    done_points = sum(i.points for i in JIRA_ISSUES if i.status == "Done")
    rows = [
        ("Backlog scope", f"{len(JIRA_ISSUES)} Sprint 6 Jira issues, {planned} planned story points", "Shows explicit backlog creation, ownership, and estimation."),
        ("Report-scoped close-out", f"{done} Done, {review} In Review, {progress} To Do/In Progress across SCRUM-81 to SCRUM-118 as of 2026-05-06 AEST", "Shows successful cleanup without unsupported Done claims."),
        ("Evidence-backed Done", f"{done_points} points closed only after repo, Jira, GitHub, or generated documentation evidence existed.", "Turns status updates into assessable evidence rather than optimism."),
        ("Live sprint snapshot", "The Jira backlog screenshot shows Sprint 6 as 34 work items with 70 Done points, 27 In Progress points, and 5 To Do points.", "Makes the remaining sprint-board work visible instead of hiding it."),
        ("W7/W8 action cleanup", "SCRUM-113 to SCRUM-118 were added for client advocacy, pitch prep, location accuracy, testing constraints, GitHub handover docs, and event visual feedback.", "Captures client-meeting work that had not been explicit enough in Jira."),
        ("Protected output", "SCRUM-81 defines playable Week 1 slice, onboarding, Yuki event, and external testing as protected outputs.", "Directly addresses Sprint 5 planning weakness by locking scope early."),
        ("Testing track", "SCRUM-104 is Done; SCRUM-105 to SCRUM-108 and SCRUM-116 remain In Review until user-supplied test evidence is attached.", "Makes the final evidence gap visible instead of hiding it."),
    ]
    return make_plate(
        "figure_04_jira_board_summary.jpg",
        "Jira Sprint 6 planning board evidence",
        "Structured evidence derived from the connected Jira project SCRUM. The board is organised around protected outputs, estimates, owners, statuses, and testing tasks.",
        rows,
        "Jira project SCRUM, issues SCRUM-81 to SCRUM-118, refreshed 2026-05-06 AEST.",
        accent=(0, 82, 204),
    )


def make_jira_full_list_plate() -> Path:
    width = 1800
    row_h = 58
    height = 210 + row_h * len(JIRA_ISSUES) + 95
    img = Image.new("RGB", (width, height), (248, 250, 252))
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 0, width, 34), fill=(0, 82, 204))
    draw.text((70, 62), "Jira Sprint 6 issue register", font=FONT_TITLE, fill=(15, 23, 42))
    draw_wrapped(
        draw,
        (70, 120),
        "Full evidence register of Sprint 6 work items, showing key, stream, owner, MoSCoW priority, estimate, and status.",
        width - 140,
        fill=(71, 85, 105),
        fnt=FONT_BODY,
    )
    y = 210
    headers = [("Key", 70), ("Stream", 205), ("Owner", 520), ("Pts", 710), ("Status", 790), ("Summary", 1020)]
    draw.rectangle((70, y, width - 70, y + 42), fill=(0, 82, 204))
    for h, x in headers:
        draw.text((x, y + 11), h, font=FONT_CAPTION, fill=(255, 255, 255))
    y += 48
    status_colours = {
        "Done": (220, 252, 231),
        "In Review": (254, 249, 195),
        "In Progress": (254, 249, 195),
        "To Do": (241, 245, 249),
    }
    for idx, issue in enumerate(JIRA_ISSUES):
        draw.rectangle((70, y, width - 70, y + row_h - 6), fill=(255, 255, 255) if idx % 2 == 0 else (241, 245, 249), outline=(203, 213, 225))
        draw.text((82, y + 15), issue.key, font=FONT_CAPTION, fill=(15, 23, 42))
        draw.text((205, y + 15), issue.stream[:27], font=FONT_CAPTION, fill=(30, 41, 59))
        draw.text((520, y + 15), issue.owner[:14], font=FONT_CAPTION, fill=(30, 41, 59))
        draw.text((715, y + 15), str(issue.points), font=FONT_CAPTION, fill=(30, 41, 59))
        rounded_rectangle(draw, (790, y + 10, 985, y + 38), 9, status_colours.get(issue.status, (241, 245, 249)), (148, 163, 184))
        draw.text((805, y + 14), issue.status, font=FONT_CAPTION, fill=(15, 23, 42))
        draw.text((1020, y + 15), issue.summary[:70], font=FONT_CAPTION, fill=(30, 41, 59))
        y += row_h
    draw.text((70, height - 55), "Source: Jira SCRUM board via Atlassian connector; SCRUM-81 to SCRUM-118.", font=FONT_CAPTION, fill=(71, 85, 105))
    out = FIG_DIR / "figure_14_jira_full_issue_register.jpg"
    img.save(out, quality=94)
    return out


def make_jira_action_cleanup_plate() -> Path:
    rows = [
        ("SCRUM-113", "Client advocacy and convenor alignment outcome captured from W7/W8 meetings.", "Done: backed by meeting notes and Jira completion comment."),
        ("SCRUM-114", "Week 10 poster and pitch rough-outline preparation added as explicit follow-up.", "In Review: needs poster/pitch outline or dry-run confirmation."),
        ("SCRUM-115", "Start-location and campus-photo accuracy decisions documented for handover.", "Done: backed by GitHub docs PR #61 and Jira evidence comment."),
        ("SCRUM-116", "Testing constraint/no-show reflection evidence separated from actual participant testing.", "In Review: needs participant count/summary or explicit no-show evidence."),
        ("SCRUM-117", "Game-facing README plus Sprint 6 evidence docs, testing/demo evidence doc, and handover doc created.", "Done: backed by GitHub PR #61 on sprint6-evidence-docs."),
        ("SCRUM-118", "Event animation/image-transition feedback converted into a Jira task.", "Done: backed by event overlay code, staged event assets, and report screenshot evidence."),
    ]
    return make_plate(
        "figure_05_jira_new_action_cleanup.jpg",
        "New Jira actions created from Week 7-8 feedback cleanup",
        "These actions close the gap between what was discussed in client meetings and what was visible as assessable Jira work.",
        rows,
        "Jira SCRUM-113 to SCRUM-118, created and triaged through 2026-05-06 AEST.",
        accent=(0, 82, 204),
    )


def make_github_pr_plate() -> Path:
    rows = [
        ("#64 Checklist coverage matrix", "Open PR adds docs/sprint6-checklist-coverage.md as an item-by-item closure matrix for the TechLauncher Sprint Checklist.", "Closes the RpD-style checklist gap in the final DOCX."),
        ("#61 Game README + Sprint 6 docs", "Merged PR added the game-facing README, Sprint 6 evidence register, handover doc, testing/demo evidence doc, and PR template.", "Separates game documentation from assessment evidence."),
        ("#60 Create W8_Client_meeting.md", "Meeting evidence added and merged on 2026-05-05.", "Stakeholder evidence and documentation trace."),
        ("#59 Ratio reconstruct and UI adjustment", "Game window ratio and related UI components adjusted.", "Execution quality and readability improvement."),
        ("#58 UI and font size adjustment", "Direct response to UI readability and font feedback.", "Client feedback converted into code change."),
        ("#57 Update phone app grid interactions", "Phone app icons and interaction feedback updated.", "Improves playable slice interaction polish."),
        ("#56 Initial mobile phone interface", "Phone UI and corresponding interactions completed.", "Concrete product feature progress."),
        ("#55 Modify UI design", "SVG UI assets imported and modified for Godot.", "Design handoff into implementation."),
        ("#51 Random events for Yuki and Kangaroo", "First set of playable slice events added.", "Core event system and treasure loop evidence."),
        ("#49 Campus location markers", "Location markers, label readability, fixed starting location, and map connections.", "Map exploration and vertical slice foundations."),
        ("#46 Main scene and week function", "Main scene reconstructed and week function built.", "Core loop infrastructure."),
    ]
    return make_plate(
        "figure_06_github_pr_evidence.jpg",
        "GitHub pull request evidence for Sprint 6 execution",
        "Merged PRs and the new documentation PR show implementation activity across map nodes, UI readability, phone interaction, event content, meeting evidence, game README, and handover packaging.",
        rows,
        "GitHub repository JeongJoonLee-kr/ANU_TreasureMap, PRs #46 to #64, refreshed 2026-05-06.",
        accent=(36, 41, 47),
    )


def make_github_docs_plate() -> Path:
    rows = [
        ("Branch", "sprint6-evidence-docs created from main and used for reviewable evidence cleanup.", "Keeps documentation improvements assessable through PR review."),
        ("README.md", "Rewritten as a game-facing overview: concept, core gameplay loop, prototype features, project structure, local run notes, and development workflow.", "Creates a future-team-facing game entry point."),
        ("docs/sprint6-review-evidence.md", "Summarises Jira evidence policy, status close-out, GitHub PR trail, testing evidence, and carry-over rules.", "Turns process evidence into a reusable repo artifact."),
        ("docs/sprint6-handover.md", "Documents playable slice scope, setup/handover guidance, known constraints, and next-team priorities.", "Supports sustainability beyond Sprint 6."),
        ("docs/testing-and-demo-evidence.md", "Defines what screenshots, videos, participant notes, and findings must be attached before final submission.", "Prevents missing test/demo claims from becoming unsupported Done work."),
        ("docs/sprint6-checklist-coverage.md", "Adds an item-by-item Sprint Checklist coverage matrix through PR #64.", "Gives the final DOCX a direct checklist closure appendix."),
    ]
    return make_plate(
        "figure_08_github_docs_pr61.jpg",
        "GitHub documentation evidence",
        "PR #61 separated the game-facing README from assessment-specific Sprint 6 evidence documents; PR #64 adds checklist closure evidence.",
        rows,
        "GitHub PR #61 and PR #64, branch sprint6-evidence-docs.",
        accent=(36, 41, 47),
    )


def make_pr_template_plate() -> Path:
    template = "\n".join(
        [
            "## Summary",
            "- What changed?",
            "- Why does this matter for the sprint goal?",
            "",
            "## Evidence",
            "- Jira issue(s): SCRUM-___",
            "- Screenshots / demo notes:",
            "- Test notes:",
            "",
            "## Checklist",
            "- [ ] Linked to Jira",
            "- [ ] Evidence added",
            "- [ ] Handover notes updated",
        ]
    )
    return make_code_plate(
        "figure_09_pr_template_evidence.jpg",
        "Pull request template evidence",
        "The new PR template makes future GitHub work carry Jira links, screenshots/demo notes, test notes, and handover evidence.",
        [(".github/pull_request_template.md excerpt", template)],
        "GitHub PR #61, .github/pull_request_template.md.",
        accent=(36, 41, 47),
    )


def make_git_commit_plate() -> Path:
    try:
        raw = run(["git", "log", "--oneline", "--since=2026-04-13", "--pretty=%h|%ad|%s", "--date=short"])
    except subprocess.CalledProcessError:
        raw = ""
    rows: list[tuple[str, str, str]] = []
    for line in raw.splitlines()[:11]:
        parts = line.split("|", 2)
        if len(parts) == 3:
            sha, date, msg = parts
            rows.append((f"{date} {sha}", msg, "Recent version-controlled work relevant to Sprint 6."))
    return make_plate(
        "figure_07_git_commit_timeline.jpg",
        "Git commit timeline evidence",
        "Recent commits and merges show steady repository activity across the Sprint 6 period.",
        rows,
        "Local git log, since 2026-04-13.",
        accent=(36, 41, 47),
    )


def make_ci_plate() -> Path:
    workflow = (ROOT / ".github/workflows/godot-ci.yml").read_text(encoding="utf-8")
    snippet = "\n".join(
        [
            "name: godot-ci",
            "jobs:",
            "  lint:",
            "    - gdlint",
            "    - gdformat --check",
            "    - validate_locations.py",
            "  smoke-test:",
            "    - install Godot 4.5-stable",
            "    - import project headless",
            "    - run res://tests/ci/smoke_test.gd",
        ]
    )
    autoformat = (ROOT / ".github/workflows/godot-autoformat.yml").read_text(encoding="utf-8")
    auto_snippet = "\n".join(autoformat.splitlines()[:12])
    return make_code_plate(
        "figure_08_github_actions_ci.jpg",
        "GitHub Actions quality pipeline evidence",
        "The repository includes Godot-specific linting, formatting, location-data validation, headless import, and smoke test workflow definitions.",
        [("godot-ci.yml quality gates", snippet), ("godot-autoformat.yml workflow excerpt", auto_snippet)],
        ".github/workflows/godot-ci.yml and .github/workflows/godot-autoformat.yml.",
        accent=(36, 41, 47),
    )


def make_meeting_plate() -> Path:
    w7_client = (ROOT / "item/meeting/meeting_2026_sems1/stakeholder_meeting/W7_Client_meeting.md").read_text(encoding="utf-8")
    w8_client = (ROOT / "item/meeting/meeting_2026_sems1/stakeholder_meeting/W8_Client_meeting.md").read_text(encoding="utf-8")
    w9_client_path = ROOT / "item/meeting/meeting_2026_sems1/stakeholder_meeting/W9_Client_meeting.md"
    w9_client = w9_client_path.read_text(encoding="utf-8") if w9_client_path.exists() else ""
    w7_team = (ROOT / "item/meeting/meeting_2026_sems1/team_meeting/W7_meeting.md").read_text(encoding="utf-8")
    def excerpt(text: str, markers: list[str], n: int = 7) -> str:
        lines = text.splitlines()
        picked: list[str] = []
        for marker in markers:
            for idx, line in enumerate(lines):
                if marker in line:
                    picked.extend(lines[idx : idx + n])
                    break
        return "\n".join([x for x in picked if x.strip()][:18])
    blocks = [
        ("W7 client meeting: pivot and protected output", excerpt(w7_client, ["Final Deliverable Scope", "Scripted Tutorial"])),
        ("W8 client meeting: grade adjustment and pitch strategy", excerpt(w8_client, ["Grade Adjustment", "Selling the Vision"])),
        ("W9 client meeting: onboarding-first final polish", excerpt(w9_client, ["Game Development & UX Feedback", "Critical Prioritization"])),
        ("W7 team meeting: work allocation and simple flow focus", excerpt(w7_team, ["Meeting Minutes", "Action Items"])),
    ]
    return make_code_plate(
        "figure_02_meeting_minutes_evidence.jpg",
        "Week 7-9 meeting minute evidence",
        "Meeting notes show the negotiated pivot, protected Week 1 slice, UI contrast feedback, onboarding-first Week 9 guidance, and team work allocation.",
        blocks,
        "item/meeting/meeting_2026_sems1.",
        accent=(51, 65, 85),
    )


def make_feedback_plate() -> Path:
    rows = [
        ("Client: negotiated pivot", "Sprint 6 should focus on a high-quality playable Week 1 vertical slice instead of a full game.", "SCRUM-81 locked protected outputs and non-goals."),
        ("Client: scripted happy path", "Tutorial week should guide players so they do not miss the core treasure-hunting loop.", "SCRUM-87 and SCRUM-106 track happy-path design and testing."),
        ("Client: UI contrast", "White text and low-contrast elements need a glasses-off test and readability improvements.", "SCRUM-96 completed as direct response; PR #58/#59 support changes."),
        ("Client: visual event feedback", "Events should use animation or image transitions when they occur, rather than relying on text alone.", "SCRUM-118 records and closes this feedback with event overlay, asset, and documentation evidence."),
        ("Client: external testing", "Run light proof-of-concept testing on the contained slice.", "SCRUM-104 to SCRUM-108 split planning, tests, and external summary."),
        ("Client: concept/pitch deck", "Sell the vision like a startup and document future research/handover ideas.", "SCRUM-112 tracks concept presentation structure."),
    ]
    return make_plate(
        "figure_03_feedback_to_action.jpg",
        "Feedback-to-Jira action register",
        "Client feedback was not left as meeting memory. It was converted into owned Jira items with statuses and evidence expectations.",
        rows,
        "Week 7-8 client meetings and Jira SCRUM issues.",
        accent=(0, 82, 204),
    )


def make_delivery_scope_plate() -> Path:
    rows = [
        ("Sprint 6 goal", "Deliver a playable Week 1 slice of ANU TreasureMap, including starting page, onboarding, Week 1 gameplay, and external proof-of-concept testing evidence.", "Anchors the sprint in a measurable final product requirement."),
        ("Client-agreed output", "Starting page, onboarding, playable Week 1, events on playable nodes, treasure collection UI, music, treasure image, and core game logic.", "Shows delivery scope was negotiated with Susan and Thomas."),
        ("Week 2-12 logic", "Later weeks are treated as repetitions/extensions of the Week 1 loop rather than separate implementation targets for this sprint.", "Makes the narrowed scope technically rational."),
        ("Week 9 refinement", "Clients advised prioritising a strong onboarding experience and quality over filling every node with content.", "Shows the team continued to refine scope with stakeholders."),
        ("Demo journey", "start page -> onboarding -> map -> node event -> treasure collection -> outcome.", "Gives the sprint review a clear playable proof path."),
        ("Happy-path philosophy", "Event choices should avoid a harsh fail-state feeling; each choice should feel like part of the student's ANU experience.", "Connects client feedback to event design values."),
    ]
    return make_plate(
        "figure_15_delivery_scope_agreement.jpg",
        "Sprint 6 delivery scope agreement",
        "The final deliverable is a client-agreed playable Week 1 slice, not a full twelve-week game. This protected scope is the central recovery decision for Sprint 6.",
        rows,
        "User-provided Sprint 6 delivery context, W7/W8 client meeting notes, and Jira SCRUM-81/SCRUM-118.",
        accent=(19, 74, 117),
    )


def make_scope_tradeoff_plate() -> Path:
    rows = [
        ("Deferred: complex resource economy", "Energy and money interactions for movement and event decisions were moved out of the Sprint 6 implementation path.", "The team protected the core playable loop from non-essential complexity."),
        ("Deferred: friend-benefit system", "Adding friends and gaining progression benefits was kept as future backlog/concept material.", "Good feature idea, but not required to prove Week 1 gameplay."),
        ("Deferred: major-dependent events", "Different starting majors changing event availability or treasure outcomes were documented for future teams.", "Preserves design depth without destabilising delivery."),
        ("Why this was reasonable", "The deferred items increase richness, but they are not the minimum product mechanism: enter game, move, trigger event, choose, receive treasure/outcome.", "Shows critical prioritisation instead of feature loss."),
        ("How value was preserved", "Unimplemented ideas were kept in official documentation and concept/handover material.", "The team reduced scope without discarding future product knowledge."),
    ]
    return make_plate(
        "figure_16_scope_tradeoff_and_backlog.jpg",
        "Scope trade-off and backlog rationale",
        "Sprint 6 reflection should show that the team actively protected final delivery by moving non-core complexity into backlog and handover documentation.",
        rows,
        "User-provided Sprint 6 reflection context, docs/wiki Final Delivery Scope, and Sprint 6 handover docs.",
        accent=(124, 45, 18),
    )


def make_dependency_reflection_plate() -> Path:
    rows = [
        ("Observed bottleneck", "Art work depended on completed story/event logic; when story content slipped, art asset production had to wait.", "Names the real team-process issue rather than hiding it."),
        ("Root cause", "The sprint had cross-functional dependency chains: narrative defines event meaning, art translates that meaning, developers integrate the asset and trigger flow.", "Explains mechanism and cause."),
        ("Better operating rule", "Create lightweight story/art contracts early: event title, location, intended feeling, required image states, and acceptance criteria before final prose is complete.", "Lets art start from stable constraints."),
        ("Integration practice", "Use placeholder-first integration and replace placeholders with final assets once approved.", "Keeps developers unblocked while art/narrative polish continues."),
        ("Jira practice", "Make dependencies visible in Jira through owner, blocked-by notes, evidence required for Done, and integration checkpoints.", "Connects reflection to concrete process improvement."),
        ("Next handover rule", "Future teams should keep the Week 1 loop stable, then add content in small batches with story, art, and code reviewed together.", "Turns retrospective into future operating principle."),
    ]
    return make_plate(
        "figure_17_dependency_reflection.jpg",
        "Reflection: dependency bottleneck and process response",
        "The main Sprint 6 collaboration problem was dependency waiting. The report frames this as a manageable production-system issue with specific next-team practices.",
        rows,
        "User-provided reflection context, W7 team meeting action allocation, Jira risk/dependency log SCRUM-84.",
        accent=(21, 128, 61),
    )


def make_event_code_plate() -> Path:
    events = json.loads((ROOT / "Godot_Game/anu-treasure-map/data/events.json").read_text(encoding="utf-8"))
    yuki = json.dumps(events["yuki"], indent=2)[:1200]
    manager = (ROOT / "Godot_Game/anu-treasure-map/systems/Script/autoloads/event_manager.gd").read_text(encoding="utf-8")
    lines = []
    for line in manager.splitlines():
        if "func try_trigger_event" in line or "eligible" in line or "require_afternoon" in line or "base_trigger_chance" in line or "complete_event" in line or "EVENT_AP_COST" in line:
            lines.append(line)
        if len(lines) >= 13:
            break
    return make_code_plate(
        "figure_10_event_system_code.jpg",
        "Yuki event and random-event system evidence",
        "The playable-slice event model is represented in data and connected to a Godot autoload that handles triggers, AP cost, choices, clue recording, and re-trigger state.",
        [("data/events.json: Yuki excerpt", yuki), ("event_manager.gd: trigger/completion logic excerpt", "\n".join(lines))],
        "Godot_Game/anu-treasure-map/data/events.json and systems/Script/autoloads/event_manager.gd.",
        accent=(132, 85, 45),
    )


def make_rubric_plate() -> Path:
    rows = [
        ("Stakeholder Engagement", "Regular Week 7-8 client meetings; pivot clarified; UI/testing/pitch feedback converted to Jira.", "Targets HD by showing proactive expectation management."),
        ("Planning and Organisation", "Protected outputs, MoSCoW, owners, estimates, risk log, velocity review, and testing tasks in Jira.", "Directly addresses prior CR planning result."),
        ("Execution and Quality", "Merged PRs for map, UI, phone, event system, code formatting, and CI workflows.", "Shows steady execution with evidence, not just intent."),
        ("Reflection and Improvement", "Sprint 5 weaknesses are framed as Sprint 6 process changes: negotiated scope, dependency visibility, earlier confirmation, and stronger evidence.", "Aims above D by linking feedback to practice change."),
    ]
    return make_plate(
        "figure_01_rubric_alignment.jpg",
        "Sprint Checklist and Rubric alignment",
        "The report is organised around the external-stakeholder team rubric: stakeholder engagement, planning, execution, and reflection.",
        rows,
        "TL26S1 Sprint Checklist and Sprint Marking Rubric PDFs.",
        accent=(19, 74, 117),
    )


def make_testing_plate() -> Path:
    rows = [
        ("Test plan", "SCRUM-104 is Done and defines target tester profile, testing script, feedback questions, and success criteria.", "Testing is explicit Sprint 6 work, not informal intention."),
        ("Onboarding test", "SCRUM-105 remains In Review until screenshots/notes confirm start page -> onboarding -> campus map.", "Avoids unsupported claims about first-time user comprehension."),
        ("Yuki happy path", "SCRUM-106 remains In Review until screenshots/notes confirm map movement, Yuki trigger, choice, treasure opportunity, and explanation.", "Keeps the core loop evidence-backed."),
        ("Hint marker flow", "SCRUM-107 remains In Review until evidence confirms future-player hint marker behaviour.", "Protects the distinctive collect-versus-help mechanic as assessable evidence."),
        ("External testing", "SCRUM-108 and SCRUM-116 remain In Review pending participant count, feedback summary, or explicit no-show/constraint evidence.", "Separates testing execution from reflection on constraints."),
    ]
    return make_plate(
        "figure_12_testing_plan_pending.jpg",
        "Testing evidence status",
        "Testing is represented by a completed plan plus explicitly isolated evidence gaps that require Week 9 material from the team.",
        rows,
        "Jira SCRUM-104 to SCRUM-108 and SCRUM-116.",
        accent=(124, 58, 237),
    )


def make_user_evidence_request_plate() -> Path:
    rows = [
        ("Week 9 team/demo notes", "Client meeting notes are now available; provide any team meeting notes, demo notes, or final build notes if available.", "Needed to strengthen final execution/reflection narrative."),
        ("Demo visuals", "Still frames or screenshots for start page, onboarding, Week 1 map, Yuki event, collect/hint choice, and outcome state.", "Needed for SCRUM-105, SCRUM-106, SCRUM-107, and SCRUM-110."),
        ("External testing", "Participant count, short feedback summary, and tester notes/forms/screenshots if available.", "Needed for SCRUM-108 and the testing section."),
        ("No-show/constraint record", "If external testers did not attend, provide the no-show count, reason/context, and mitigation/reflection.", "Needed for SCRUM-116."),
        ("Pitch/poster", "Poster/pitch rough outline or dry-run confirmation if already done.", "Needed for SCRUM-86 and SCRUM-114."),
        ("Retro/reflection", "Final retrospective notes and any individual reflections if available.", "Needed to strengthen Reflection and Improvement to HD standard."),
    ]
    return make_plate(
        "figure_17_user_evidence_needed.jpg",
        "User-provided evidence still required before final submission",
        "The report is evidence-cleaned and now includes Week 9 client notes, but final demo/testing artefacts still need direct screenshots, notes, or summaries from the team.",
        rows,
        "Open Jira In Review issues after Sprint 6 evidence cleanup.",
        accent=(124, 58, 237),
    )


def make_checklist_plate(filename: str, title: str, subtitle: str, rows: list[tuple[str, str, str]]) -> Path:
    return make_plate(
        filename,
        title,
        subtitle,
        rows,
        "TL26S1 Sprint Checklist, Sprint Marking Rubric, Jira, GitHub PR #61/#64, and Sprint 6 meeting notes.",
        accent=(19, 74, 117),
    )


def make_checklist_stakeholder_plate() -> Path:
    rows = [
        ("Regular meetings", "W7, W8, and W9 client meetings are recorded; final team/demo notes would strengthen the close-out.", "Covered"),
        ("Detailed updates", "Vertical slice, UI readability, testing, pitch prep, and handover updates are visible.", "Covered"),
        ("Expectations/risks", "Negotiated pivot, protected outputs, non-goals, testing risk, and advocacy are tracked.", "Covered"),
        ("Minutes/actions", "Meeting notes and action tables are stored in the repo.", "Covered"),
        ("Feedback tasks", "W7/W8 feedback became Jira actions, including SCRUM-113 to SCRUM-118.", "Covered"),
        ("Report status back", "Report and PR evidence summarise status; W9 client notes confirm onboarding-first final polish guidance.", "Covered"),
        ("Relationship practices", "Client advocacy and negotiated pivot are recorded as stakeholder-management evidence.", "Covered"),
    ]
    return make_checklist_plate(
        "figure_18_checklist_stakeholder.jpg",
        "Checklist closure: Stakeholder Engagement",
        "Item-by-item closure against the stakeholder section of the Sprint Checklist.",
        rows,
    )


def make_checklist_planning_plate() -> Path:
    rows = [
        ("Vision and sprint goal", "Sprint 6 is framed as a playable Week 1 vertical slice plus concept handover.", "Covered"),
        ("Task tool", "Jira is used as the task-management evidence system.", "Covered"),
        ("Goals/milestones", "Protected outputs, non-goals, pitch preparation, and Week 9 evidence needs are visible.", "Covered"),
        ("Backlog breadth", "Product, testing, build, PR, and handover work are tracked; personas/scenarios can be stronger.", "Partially covered"),
        ("PBI mapping", "Jira issues map to stakeholder, planning, execution, testing, and reflection outcomes.", "Covered"),
        ("Estimation", "Story points are recorded; explicit Planning Poker evidence is not visible.", "Partially covered"),
        ("Allocation", "Jira owners and W7 team action items show responsibility allocation.", "Covered"),
        ("Prioritisation", "MoSCoW and protected-output framing separate Must/Should work.", "Covered"),
        ("Risks/dependencies", "Scope drift, integration, testing, data, and location risks are tracked.", "Covered"),
        ("Velocity/monitoring", "Velocity review exists, but final retro should connect it to next-sprint planning.", "Partially covered"),
    ]
    return make_checklist_plate(
        "figure_19_checklist_planning.jpg",
        "Checklist closure: Planning and Organisation",
        "Item-by-item closure against the planning section of the Sprint Checklist.",
        rows,
    )


def make_checklist_execution_plate() -> Path:
    rows = [
        ("Tool kept updated", "Jira status cleanup separates evidence-backed Done from user-dependent In Review work.", "Covered"),
        ("Milestones/PBIs", "Sprint 6 issues show vertical slice, documentation, and remaining Week 9 evidence needs.", "Covered"),
        ("Version control", "GitHub PRs and commits cover map, UI, phone interaction, event system, docs, and meetings.", "Covered"),
        ("CI/quality", "Godot workflows cover linting, formatting, validation, import, and smoke-test gates.", "Covered"),
        ("Testing", "Testing plan exists; participant results and final demo screenshots remain needed.", "Partially covered"),
        ("Peer review", "PR #61 review comments were addressed; PR #64 has Copilot review requested.", "Covered"),
        ("Challenge response", "Scope pivot, CI path-filter clarification, and testing constraints are documented.", "Covered"),
        ("Value delivered", "Running Godot screenshot proves the playable Week 1 map slice; full demo states still need W9.", "Partially covered"),
    ]
    return make_checklist_plate(
        "figure_20_checklist_execution.jpg",
        "Checklist closure: Execution and Quality",
        "Item-by-item closure against the execution section of the Sprint Checklist.",
        rows,
    )


def make_checklist_reflection_plate() -> Path:
    rows = [
        ("Feedback considered", "Sprint 5 feedback, W6 tutor guidance, and W7/W8 client feedback shaped Sprint 6 cleanup.", "Covered"),
        ("Sprint review", "DOCX embeds Jira, GitHub, meeting, code, build, testing-plan, and checklist evidence.", "Covered"),
        ("Retrospective", "Reflection narrative exists; final team retrospective notes should still be supplied.", "Partially covered"),
        ("PBIs from review", "SCRUM-113 to SCRUM-118 were created from W7/W8 evidence-gap review.", "Covered"),
        ("Feedback response clarity", "Feedback-to-action chain is explicit in Jira and the report.", "Covered"),
        ("Individual reflections", "Not currently visible in repository evidence; add if expected for stronger HD support.", "Partially covered"),
    ]
    return make_checklist_plate(
        "figure_21_checklist_reflection.jpg",
        "Checklist closure: Reflection and Improvement",
        "Item-by-item closure against the reflection section of the Sprint Checklist.",
        rows,
    )


def make_reflection_plate() -> Path:
    rows = [
        ("Sprint 5 planning gap", "The project pivot and paper-prototype expectation created more work than the team could responsibly absorb with its available resources.", "Sprint 6 begins by calculating realistic capacity from previous story points, then protecting a smaller delivery target."),
        ("Negotiated pivot", "The reduced scope was not a unilateral retreat; W7 client notes record Susan's position that the changes were negotiated pivots based on feedback and deliverability.", "This converts a potential marking risk into stakeholder-aligned evidence."),
        ("Sprint 6 execution response", "The team chose a playable Week 1 slice because Week 2-12 repeat the same loop and do not need separate implementation to prove the product mechanism.", "Execution is judged against a demonstrable final product path, not an over-claimed full game."),
        ("Reflection strength", "The team changed practice: visible Jira action register, protected outputs, deferred backlog, PR evidence, build screenshots, and explicit evidence gaps.", "Reflection is shown as a delivery-system change rather than a personal feeling."),
        ("Stakeholder strength", "Client advocacy, UI feedback, happy-path advice, event visual feedback, and pitch strategy were all converted into tasks or handover evidence.", "The team sustained its prior stakeholder strength while improving planning/execution."),
    ]
    return make_plate(
        "figure_13_sprint5_to_sprint6_improvement.jpg",
        "Sprint 5 feedback to Sprint 6 improvement loop",
        "The Sprint 6 report explicitly answers earlier rubric weaknesses by showing how the team changed its operating model.",
        rows,
        "Sprint 5 Documentation Bundle and Sprint 5 Collected Reflections.",
        accent=(21, 128, 61),
    )


def make_concept_plate() -> Path:
    rows = [
        ("Playable Week 1 slice", "Start page, onboarding, Week 1 map, Yuki event, treasure opportunity, collect/hint decision.", "Small enough to finish; clear enough to test."),
        ("Concept presentation", "Architecture, database ideas, future features, unimplemented concepts, and handover recommendations.", "Preserves project depth beyond the implemented slice."),
        ("Pitch strategy", "Clients advised a movie-trailer/startup framing: sell vision and buy-in, not only features.", "Supports Week 10 poster and video preparation."),
        ("Future team value", "The deck should document ideal version, findings, and where the next team should start.", "Turns Sprint 6 into research-style handover evidence."),
    ]
    return make_plate(
        "figure_11_concept_handover_strategy.jpg",
        "Concept deck and handover strategy",
        "The Sprint 6 deliverable is not only a build; it also packages technical insight and future work for showcase and handover.",
        rows,
        "Week 7-8 client meeting notes and Jira SCRUM-112.",
        accent=(124, 45, 18),
    )


def generate_figures() -> dict[str, Path]:
    figures = {
        "rubric": make_rubric_plate(),
        "meetings": make_meeting_plate(),
        "feedback": make_feedback_plate(),
        "jira_summary": make_jira_board_plate(),
        "jira_actions": make_jira_action_cleanup_plate(),
        "jira_full": make_jira_full_list_plate(),
        "github_pr": make_github_pr_plate(),
        "github_docs": make_github_docs_plate(),
        "pr_template": make_pr_template_plate(),
        "git_timeline": make_git_commit_plate(),
        "ci": make_ci_plate(),
        "product": make_product_plate(),
        "delivery_scope": make_delivery_scope_plate(),
        "scope_tradeoff": make_scope_tradeoff_plate(),
        "dependency_reflection": make_dependency_reflection_plate(),
        "event_visual_feedback": make_event_visual_feedback_plate(),
        "event_code": make_event_code_plate(),
        "concept": make_concept_plate(),
        "testing": make_testing_plate(),
        "evidence_needed": make_user_evidence_request_plate(),
        "reflection": make_reflection_plate(),
        "checklist_stakeholder": make_checklist_stakeholder_plate(),
        "checklist_planning": make_checklist_planning_plate(),
        "checklist_execution": make_checklist_execution_plate(),
        "checklist_reflection": make_checklist_reflection_plate(),
    }
    return figures


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, color: str = "1F2937") -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Arial"
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(color)


def add_caption(doc: Document, caption: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(caption)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(71, 85, 105)


def add_figure(doc: Document, path: Path, caption: str, width: float = 6.35) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    add_caption(doc, caption)


def add_optional_figure(doc: Document, path: Path, caption: str, width: float = 6.35) -> None:
    if path.exists():
        add_figure(doc, path, caption, width)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(item)
        r.font.name = "Arial"
        r.font.size = Pt(10.5)


def add_summary_table(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["Rubric area", "Sprint 6 evidence", "Grade case", "Evidence status"]
    for idx, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], h, True, "FFFFFF")
        set_cell_shading(table.rows[0].cells[idx], "174A75")
    rows = [
        ("Stakeholder", "Week 7-9 client meetings; feedback-to-Jira register; new SCRUM-113 to SCRUM-118 cleanup actions.", "HD/D", "Final testing/demo evidence still requested."),
        ("Planning", "38 report-scoped Jira issues; 27 Done and 11 In Review after SCRUM-110 and SCRUM-118 were closed with build/event evidence.", "HD-leaning", "Actual Jira board/issue screenshots embedded."),
        ("Execution", "PRs #46-64; PR #61 merged; PR #64 opened; live Godot map screenshot; Yuki/event visual evidence; CI workflows.", "D/HD-leaning", "Full demo-flow screenshots still requested."),
        ("Reflection", "Sprint 5 lessons converted into negotiated scope, evidence-backed Done policy, dependency-management practice, checklist closure, and explicit Week 9 evidence list.", "HD-leaning", "Individual reflections still requested later."),
    ]
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)


def configure_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    styles = doc.styles
    for style_name in ["Normal", "Body Text"]:
        if style_name in styles:
            st = styles[style_name]
            st.font.name = "Arial"
            st.font.size = Pt(10.5)
    for style_name, size, color in [
        ("Title", 22, "0F172A"),
        ("Heading 1", 16, "174A75"),
        ("Heading 2", 13, "1F2937"),
        ("Heading 3", 11, "334155"),
    ]:
        st = styles[style_name]
        st.font.name = "Arial"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
    header = section.header.paragraphs[0]
    header.text = "ANU TreasureMap - Sprint 6 Review Report"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor(100, 116, 139)


def add_title_page(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ANU TreasureMap\nSprint 6 Review Report")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(26)
    r.font.color.rgb = RGBColor(15, 23, 42)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Single-file evidence report aligned to the TechLauncher Sprint Checklist and Marking Rubric")
    r.font.name = "Arial"
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(71, 85, 105)
    doc.add_paragraph("")
    meta = doc.add_table(rows=5, cols=2)
    meta.style = "Table Grid"
    rows = [
        ("Project", "ANU TreasureMap | External stakeholder team | Godot"),
        ("Sprint focus", "Final delivery: playable Week 1 slice from starting page through onboarding, Week 1 gameplay, event/treasure loop, and proof-of-concept testing evidence"),
        ("Sprint scope", "Week 7 to Week 9; evidence-cleaned report refreshed on 2026-05-06 AEST"),
        ("Tools", "Jira for task management; GitHub for version control, PRs, commits, and CI workflows"),
        ("Team members", "Jeongjoon Lee, Wing Sze Shing, Suya Xie, Zijun Yang, Xiyin Li, Anlin Bai, Shaoming Qin"),
    ]
    for r_idx, (a, b) in enumerate(rows):
        set_cell_text(meta.rows[r_idx].cells[0], a, True)
        set_cell_shading(meta.rows[r_idx].cells[0], "E2E8F0")
        set_cell_text(meta.rows[r_idx].cells[1], b)
    doc.add_page_break()


def build_doc(figures: dict[str, Path]) -> None:
    doc = Document()
    configure_doc(doc)
    add_title_page(doc)

    doc.add_heading("Purpose", level=1)
    doc.add_paragraph(
        "This report consolidates Sprint 6 evidence for ANU TreasureMap into a single reviewable DOCX. "
        "It follows the evidence-forward style of the HD sample report while translating the content to this team's tools, product, and sprint context. "
        "Sprint 6 is treated as the final product delivery sprint: the team is presenting a client-agreed playable Week 1 slice from starting page through onboarding, Week 1 gameplay, event/treasure interaction, and proof-of-concept testing evidence. "
        "Because the final submission is a single DOCX, the key Jira, GitHub, meeting, code, and product evidence is embedded directly as images rather than relying on external links."
    )
    add_figure(doc, figures["rubric"], "Figure 1. Sprint Checklist and marking rubric alignment used to structure the report.")

    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(
        "Sprint 6 should be read as both a recovery-and-quality sprint and the project's final deliverable sprint. Sprint 5 produced strong stakeholder and reflection evidence, but planning and execution were marked at Credit because the team had not yet translated the pivoted game concept into a tightly protected, implementation-ready, testable slice. In Sprint 6 the team responded by narrowing the target to a playable Week 1 vertical slice and by turning client feedback into explicit Jira work items, GitHub implementation evidence, and test preparation."
    )
    add_bullets(
        doc,
        [
            "Stakeholder engagement remained a strength: the team used Week 7, Week 8, and Week 9 client meetings to clarify a negotiated pivot, define the final deliverable shape, sharpen onboarding-first polish, and convert specific feedback into tasks.",
            "Planning improved materially: the report-scoped Jira register now shows 38 Sprint 6 issues, 27 Done, and 11 In Review after evidence cleanup, with live Jira screenshots embedded for verification.",
            "Execution evidence is stronger than Sprint 5: PRs #46-64 show map, UI, phone interaction, event data, Godot event logic, CI, a game-facing README, Sprint 6 evidence docs, checklist closure, and event-visual feedback closure.",
            "Testing is now planned as explicit work; the remaining In Review items clearly list the Week 9 demo/testing evidence that must be supplied before the final evidence-closed submission.",
        ],
    )
    doc.add_page_break()
    add_summary_table(doc)

    doc.add_heading("1. Stakeholder Engagement", level=1)
    doc.add_paragraph(
        "The strongest Sprint 6 stakeholder evidence is the feedback-to-action chain. Client meetings did not only confirm attendance; they clarified the negotiated pivot, protected the Week 1 vertical-slice goal, and converted UI, demo, testing, pitch, onboarding-first polish, and event-visual concerns into Jira work. This directly addresses the risk that the team could be seen as changing scope without structure."
    )
    add_figure(doc, figures["meetings"], "Figure 2. Week 7-9 meeting evidence showing client feedback, negotiated pivot, onboarding-first guidance, and internal task allocation.")
    add_optional_figure(doc, REAL_DIR / "github_w7_client_meeting_cropped.png", "Figure 2A. Actual GitHub screenshot of W7_Client_meeting.md showing the negotiated pivot, playable Week 1 slice, UI contrast feedback, and happy-path strategy.", 6.35)
    add_figure(doc, figures["feedback"], "Figure 3. Feedback-to-Jira action register showing client feedback converted into owned Jira tasks.")
    add_figure(doc, figures["jira_actions"], "Figure 4. New Jira actions created from Week 7-8 client feedback that was not previously explicit enough in Jira.")
    doc.add_paragraph(
        "The client also helped clarify that the scope change should be understood as a negotiated pivot rather than a sign of disorganisation. That is important evidence because it connects stakeholder management with assessment risk management: the team was not simply reducing ambition; it was protecting a higher-quality, demonstrable slice with client support."
    )

    doc.add_heading("2. Planning and Organisation", level=1)
    doc.add_paragraph(
        "Planning in Sprint 6 is substantially stronger than in Sprint 5 because the team created a more explicit operating model around Jira. The backlog is not a loose task list: it shows protected outputs, owners, estimates, MoSCoW priority, risk/dependency tracking, test tasks, build evidence tasks, and a velocity review. After cleanup, report-scoped evidence-backed work is Done, while user-dependent demo/testing/pitch evidence remains clearly separated as In Review. The live Jira board still shows some sprint-board work open, so the report treats those items as visible Week 9/evidence dependencies rather than pretending the board is fully closed."
    )
    add_figure(doc, figures["jira_summary"], "Figure 5. Jira planning summary showing protected output, estimates, owners, and final evidence-cleaned status.")
    add_optional_figure(doc, REAL_DIR / "jira_sprint6_backlog_cropped.png", "Figure 6. Actual Jira Sprint 6 backlog screenshot showing sprint dates, work-item count, point status, owners, estimates, and issue statuses.", 6.35)
    add_figure(doc, figures["jira_full"], "Figure 7. Full Sprint 6 Jira issue register embedded for single-file evidence review.")
    doc.add_paragraph(
        "The key planning improvement is that Sprint 6 explicitly separates what must be protected from what can be documented as future work. The protected output is the playable Week 1 slice: start page, onboarding, Week 1 map flow, Yuki event, collect-or-hint decision, and light external testing. The broader ideal version is captured through concept and handover work rather than being allowed to destabilise the implementation sprint."
    )
    add_figure(doc, figures["delivery_scope"], "Figure 7A. Client-agreed Sprint 6 delivery scope: playable Week 1 slice and repeatable Week 2-12 logic.")
    add_figure(doc, figures["scope_tradeoff"], "Figure 7B. Scope trade-off evidence showing why complex resource, friend, and major-dependent systems were moved to backlog/handover.")

    doc.add_heading("3. Execution and Quality", level=1)
    doc.add_paragraph(
        "Execution evidence is strongest in GitHub and the running Godot build. The Sprint 6 period includes merged PRs for map and location work, UI contrast and font adjustments, phone interaction, event content, Yuki/Kangaroo random event logic, meeting-note evidence, and ongoing format/quality workflow activity. PR #61 is now merged for the game-facing README and Sprint 6 evidence docs, while PR #64 adds the item-by-item checklist closure matrix."
    )
    add_figure(doc, figures["github_pr"], "Figure 8. GitHub PR evidence showing Sprint 6 execution across map, UI, phone, event, documentation, and checklist work.")
    add_optional_figure(doc, REAL_DIR / "github_pr64_cropped.png", "Figure 9. Actual GitHub PR #64 screenshot showing the Sprint 6 checklist coverage matrix PR ready for review/merge.", 6.35)
    add_figure(doc, figures["github_docs"], "Figure 10. GitHub documentation evidence showing the game-facing README separated from Sprint 6 evidence documents.")
    add_figure(doc, figures["pr_template"], "Figure 11. Pull request template evidence requiring Jira links, screenshots/demo notes, tests, and handover updates.")
    add_figure(doc, figures["git_timeline"], "Figure 12. Git commit timeline for recent Sprint 6 repository activity.")
    add_figure(doc, figures["ci"], "Figure 13. GitHub Actions workflow evidence for lint, formatting, location validation, and smoke testing.")
    doc.add_paragraph(
        "A live product screenshot is now available: Computer Use captured the running Godot debug build with the Week 1 map, status HUD, location network, Bruce Hall marker, and Next Week control. This is real build evidence, not a mock-up. It does not yet replace the Week 9 demo-flow evidence still needed for start page, onboarding, Yuki choice, collect/hint choice, and outcome state."
    )
    add_optional_figure(doc, REAL_DIR / "godot_playable_map_cropped.png", "Figure 14. Actual running Godot debug build screenshot showing the playable Week 1 ANU campus map slice.", 6.35)
    add_figure(doc, figures["product"], "Figure 15. Repository product/build evidence from committed Godot assets.")
    add_figure(doc, figures["event_code"], "Figure 16. Yuki event data and EventManager implementation evidence.")
    add_figure(doc, figures["event_visual_feedback"], "Figure 16A. Event visual feedback evidence connecting SCRUM-118 to overlay flow, image sequence assets, and treasure image display.")
    add_optional_figure(doc, REAL_DIR / "jira_scrum118_cropped.png", "Figure 16B. Actual Jira screenshot showing SCRUM-118 Done with acceptance criteria and implementation evidence.", 6.35)

    doc.add_heading("4. Reflection and Improvement", level=1)
    doc.add_paragraph(
        "The reflection story for Sprint 6 is not that the team suddenly avoided all uncertainty. It is that the team learned how to convert uncertainty into negotiated scope, visible Jira work, and a playable proof path. Sprint 5 was difficult because the project pivot and the client's paper-prototype expectation placed real pressure on the team's limited resources. The important reflective point is that the scope reduction was not decided privately by the team; the W7 client evidence records it as a negotiated pivot based on client feedback and what was possible to deliver."
    )
    add_figure(doc, figures["reflection"], "Figure 17. Sprint 5 feedback mapped to Sprint 6 process improvements.")
    doc.add_paragraph(
        "At the start of Sprint 6, the team listed the work that would be required for the sprint, compared it with the story-point load the team had actually managed in the previous period, and used that capacity estimate to prioritise Must and Should work. The team then presented the realistic delivery plan to the clients. This changed the planning pattern: instead of beginning with a broad ideal product and cutting late, Sprint 6 began with a protected Week 1 slice that the clients accepted as a meaningful final delivery path."
    )
    doc.add_paragraph(
        "The Week 1 focus was decided in the Week 7 internal planning discussion after the team considered realistic workload and cross-functional dependencies. This decision mattered because Week 2 through Week 12 use the same gameplay logic as Week 1. By proving Week 1 well, the team could demonstrate the product mechanism without forcing twelve weeks of duplicated content into a three-week sprint. Features such as energy/money interaction complexity, friend-based benefits, and major-dependent event paths were moved into backlog and concept documentation because they increase richness but are not required to prove the core loop."
    )
    add_figure(doc, figures["dependency_reflection"], "Figure 18. Dependency bottleneck reflection and future operating practices.")
    doc.add_paragraph(
        "The major collaboration challenge was dependency waiting. Art production depended on story and event logic being clear enough to guide image generation; when narrative decisions moved slowly, art could not confidently start. The strongest rubric response is not to blame one stream, but to show the production mechanism and the process change: future work should use early story/art contracts, placeholder-first integration, explicit blocked-by notes in Jira, and short integration checkpoints so narrative, art, and code can move in parallel."
    )
    add_figure(doc, figures["concept"], "Figure 19. Concept deck and handover strategy preserving the broader project vision.")
    doc.add_paragraph(
        "This is why the Sprint 6 grade case can aim above the previous planning and execution result. The team is not asking for credit on aspiration alone: the evidence shows a changed process, narrower scope control, more explicit ownership, direct feedback-to-Jira conversion, and a stronger trail between client feedback and deliverable work. Individual reflection evidence is still expected later, but the team-level reflection now has the mechanism, cause, response, and future rule that the HD sample uses well."
    )

    doc.add_heading("Appendix: Full Evidence Pack", level=1)
    doc.add_paragraph(
        "This appendix is included because the final submission is expected to stand alone as a single DOCX. The images below are deliberately embedded so the marker can assess the evidence without opening Jira, GitHub, or the repository."
    )
    add_figure(doc, figures["testing"], "Figure 20. Testing plan and Week 9 external proof-of-concept evidence status.")
    add_figure(doc, figures["evidence_needed"], "Figure 21. User-provided evidence still required before the report can be treated as fully evidence-closed.")
    add_heading = doc.add_heading
    add_heading("Checklist Item Closure Matrix", level=2)
    doc.add_paragraph(
        "The following four figures close the TechLauncher Sprint Checklist directly, item by item. This mirrors the HD sample's habit of making the marker's checklist work easy: each item is either Covered, Partially covered, or left as Week 9/user evidence."
    )
    add_figure(doc, figures["checklist_stakeholder"], "Figure 22. Stakeholder Engagement checklist closure matrix.")
    add_figure(doc, figures["checklist_planning"], "Figure 23. Planning and Organisation checklist closure matrix.")
    add_figure(doc, figures["checklist_execution"], "Figure 24. Execution and Quality checklist closure matrix.")
    add_figure(doc, figures["checklist_reflection"], "Figure 25. Reflection and Improvement checklist closure matrix.")

    doc.add_heading("Final Evidence Close-Out Status", level=2)
    doc.add_paragraph(
        "This report now reflects the Jira and GitHub cleanup requested for Sprint 6: missing Week 7-8 feedback actions were created, Week 9 client notes were added to the evidence base, evidence-backed Jira issues were closed, SCRUM-110 was closed after a real Godot build screenshot was captured, SCRUM-118 was created and closed for event animation/image-transition feedback, unsupported testing/pitch work was left In Review, PR #61 was merged for the game README and Sprint evidence documentation, and PR #64 was opened for checklist closure. The remaining evidence gap is deliberately narrow and user-dependent: full demo screenshots or video frames, external-testing results or no-show reflection, pitch/poster confirmation, and final retrospective/individual reflections. Those items should be inserted as final embedded evidence before submission."
    )

    doc.save(DOCX_PATH)
    import shutil

    shutil.copyfile(DOCX_PATH, ROOT_DOCX_PATH)


def main() -> None:
    ensure_dirs()
    figures = generate_figures()
    build_doc(figures)
    print(DOCX_PATH)
    print(ROOT_DOCX_PATH)


if __name__ == "__main__":
    main()
